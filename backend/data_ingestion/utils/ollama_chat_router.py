from __future__ import annotations
from data_ingestion.config.paths import PROJECT_ROOT

import asyncio
import base64
import json
import logging
import re
import sqlite3
import uuid
from pathlib import Path
from typing import Any, Callable

import requests
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from PIL import Image, ImageEnhance, ImageFilter, ImageOps


ROOT = PROJECT_ROOT
CHAT_ROOT = ROOT / "tools" / "ollama"
DB_PATH = CHAT_ROOT / "chat_history.db"
UPLOAD_DIR = CHAT_ROOT / "uploads"
ENHANCED_UPLOAD_DIR = CHAT_ROOT / "uploads_enhanced"
TEMPLATE_PATH = CHAT_ROOT / "templates" / "index.html"
MAX_TEXT = 80_000
OLLAMA_URL = "http://127.0.0.1:11434"

IMG = {"png", "jpg", "jpeg", "gif", "webp", "bmp", "tiff", "ico"}
TXT = {
    "txt", "md", "py", "js", "ts", "jsx", "tsx", "html", "css", "json", "yaml", "yml",
    "xml", "sh", "bat", "ps1", "c", "cpp", "h", "hpp", "java", "go", "rs", "rb",
    "php", "sql", "r", "swift", "kt", "scala", "lua", "toml", "ini", "conf", "log",
    "env", "dockerfile", "makefile", "tf", "proto", "vue", "svelte", "tex", "csv",
}
PDF = {"pdf"}
DOCX = {"docx", "doc"}
XLSX = {"xlsx", "xls"}

logger = logging.getLogger(__name__)


UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
ENHANCED_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def _db() -> sqlite3.Connection:
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def init_ollama_chat_db() -> None:
    with _db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS conversations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                owner TEXT NOT NULL DEFAULT 'legacy',
                title TEXT NOT NULL DEFAULT 'New Chat',
                model TEXT NOT NULL DEFAULT '',
                created_at TEXT DEFAULT (datetime('now')),
                updated_at TEXT DEFAULT (datetime('now'))
            );
            CREATE TABLE IF NOT EXISTS messages (
                id              INTEGER PRIMARY KEY AUTOINCREMENT,
                conversation_id INTEGER NOT NULL,
                role            TEXT    NOT NULL,
                content         TEXT    NOT NULL,
                attachment_name TEXT,
                attachment_type TEXT,
                attachment_path TEXT,
                created_at      TEXT DEFAULT (datetime('now')),
                FOREIGN KEY (conversation_id) REFERENCES conversations(id) ON DELETE CASCADE
            );
            """
        )
        have = {row[1] for row in conn.execute("PRAGMA table_info(messages)")}
        for col in ("attachment_name", "attachment_type", "attachment_path"):
            if col not in have:
                conn.execute(f"ALTER TABLE messages ADD COLUMN {col} TEXT")
        conversation_cols = {row[1] for row in conn.execute("PRAGMA table_info(conversations)")}
        if "owner" not in conversation_cols:
            conn.execute("ALTER TABLE conversations ADD COLUMN owner TEXT NOT NULL DEFAULT 'legacy'")


def _safe_name(name: str | None) -> str:
    value = Path(name or "file").name
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._")
    return value or "file"


def _ext(name: str | None) -> str:
    safe = _safe_name(name)
    return safe.rsplit(".", 1)[-1].lower() if "." in safe else ""


def _category(name: str | None) -> str:
    ext = _ext(name)
    if ext in IMG:
        return "image"
    if ext in TXT:
        return "text"
    if ext in PDF:
        return "pdf"
    if ext in DOCX:
        return "docx"
    if ext in XLSX:
        return "spreadsheet"
    return "binary"


async def _save_upload(file: UploadFile) -> tuple[str, str, str]:
    original = _safe_name(file.filename)
    ext = _ext(original)
    stored = f"{uuid.uuid4().hex}.{ext}" if ext else uuid.uuid4().hex
    target = UPLOAD_DIR / stored
    data = await file.read()
    target.write_bytes(data)
    return stored, original, _category(original)


def _image_to_jpeg_bytes(image: Image.Image, *, quality: int = 96) -> bytes:
    import io

    buf = io.BytesIO()
    image.convert("RGB").save(buf, format="JPEG", quality=quality, optimize=True)
    return buf.getvalue()


def _enhance_image_for_llm(path: Path, stored: str) -> bytes:
    """Return a sharpened/upscaled JPEG for vision models while keeping the original upload."""
    image = Image.open(path)
    image = ImageOps.exif_transpose(image).convert("RGB")

    width, height = image.size
    max_side = max(width, height)
    if max_side < 1800:
        scale = min(3.0, 1800 / max(1, max_side))
        image = image.resize((int(width * scale), int(height * scale)), Image.Resampling.LANCZOS)
    elif max_side > 2600:
        scale = 2600 / max_side
        image = image.resize((int(width * scale), int(height * scale)), Image.Resampling.LANCZOS)

    image = ImageOps.autocontrast(image, cutoff=0.5)
    image = ImageEnhance.Color(image).enhance(1.08)
    image = ImageEnhance.Contrast(image).enhance(1.28)
    image = ImageEnhance.Sharpness(image).enhance(2.4)
    image = image.filter(ImageFilter.UnsharpMask(radius=1.2, percent=240, threshold=2))
    image = ImageEnhance.Contrast(image).enhance(1.08)
    image = image.filter(ImageFilter.UnsharpMask(radius=0.7, percent=130, threshold=1))

    enhanced_bytes = _image_to_jpeg_bytes(image)
    enhanced_name = f"{Path(stored).stem}_enhanced_for_llm.jpg"
    (ENHANCED_UPLOAD_DIR / enhanced_name).write_bytes(enhanced_bytes)
    return enhanced_bytes


def _extract(stored: str, original: str) -> tuple[str, str | None]:
    category = _category(original)
    path = UPLOAD_DIR / stored
    if category == "image":
        try:
            image_bytes = _enhance_image_for_llm(path, stored)
        except Exception:
            image_bytes = path.read_bytes()
        return "image", base64.b64encode(image_bytes).decode("ascii")
    if category == "text":
        return "text", path.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT]
    if category == "pdf":
        try:
            import pypdf

            reader = pypdf.PdfReader(str(path))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
            return "text", text[:MAX_TEXT]
        except Exception as exc:
            return "text", f"[PDF read error: {exc}]"
    if category == "docx":
        try:
            import docx

            doc = docx.Document(str(path))
            return "text", "\n".join(p.text for p in doc.paragraphs if p.text)[:MAX_TEXT]
        except Exception as exc:
            return "text", f"[DOCX read error: {exc}]"
    if category == "spreadsheet":
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(str(path), data_only=True)
            rows: list[str] = []
            for sheet in workbook.worksheets:
                rows.append(f"=== {sheet.title} ===")
                for row in sheet.iter_rows(values_only=True):
                    rows.append("\t".join(str(value or "") for value in row))
            return "text", "\n".join(rows)[:MAX_TEXT]
        except Exception as exc:
            return "text", f"[Spreadsheet read error: {exc}]"
    return "binary", None


def _make_conversation(owner: str, model: str, title: str = "New Chat") -> int:
    with _db() as conn:
        cur = conn.execute("INSERT INTO conversations (owner,title,model) VALUES (?,?,?)", (owner, title, model))
        return int(cur.lastrowid)


def _require_owned_conversation(conn: sqlite3.Connection, conversation_id: int, owner: str) -> sqlite3.Row:
    row = conn.execute(
        "SELECT * FROM conversations WHERE id=? AND owner=?",
        (conversation_id, owner),
    ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="conversation not found")
    return row


def _save_user(
    conversation_id: int,
    text: str,
    attachment_name: str | None = None,
    attachment_type: str | None = None,
    attachment_path: str | None = None,
) -> None:
    with _db() as conn:
        conn.execute(
            "INSERT INTO messages (conversation_id,role,content,attachment_name,attachment_type,attachment_path)"
            " VALUES (?,'user',?,?,?,?)",
            (conversation_id, text, attachment_name, attachment_type, attachment_path),
        )
        count = conn.execute(
            "SELECT COUNT(*) FROM messages WHERE conversation_id=? AND role='user'",
            (conversation_id,),
        ).fetchone()[0]
        if count == 1:
            raw = text or attachment_name or "New Chat"
            title = raw[:60] + ("..." if len(raw) > 60 else "")
            conn.execute(
                "UPDATE conversations SET title=?,updated_at=datetime('now') WHERE id=?",
                (title, conversation_id),
            )
        else:
            conn.execute("UPDATE conversations SET updated_at=datetime('now') WHERE id=?", (conversation_id,))


def _save_assistant(conversation_id: int, text: str) -> None:
    with _db() as conn:
        conn.execute(
            "INSERT INTO messages (conversation_id,role,content) VALUES (?,'assistant',?)",
            (conversation_id, text),
        )
        conn.execute("UPDATE conversations SET updated_at=datetime('now') WHERE id=?", (conversation_id,))


def _load_history(conversation_id: int, keep_images: int = 6) -> list[dict[str, Any]]:
    with _db() as conn:
        rows = conn.execute(
            "SELECT * FROM messages WHERE conversation_id=? ORDER BY id",
            (conversation_id,),
        ).fetchall()
    total = len(rows)
    messages: list[dict[str, Any]] = []
    for index, row in enumerate(rows):
        content = row["content"] or ""
        message: dict[str, Any] = {"role": row["role"], "content": content}
        if row["attachment_type"] == "image" and row["attachment_path"] and (total - index) <= keep_images:
            path = UPLOAD_DIR / row["attachment_path"]
            if path.exists():
                message["content"] = content or "Analyze this image."
                message["images"] = [base64.b64encode(path.read_bytes()).decode("ascii")]
        messages.append(message)
    return messages


def _ollama_complete(model: str, messages: list[dict[str, Any]], conversation_id: int) -> dict[str, Any]:
    logger.info(
        "Ollama chat complete | model=%s | conversation_id=%s | messages=%s",
        model,
        conversation_id,
        len(messages),
    )
    try:
        response = requests.post(
            f"{OLLAMA_URL}/api/chat",
            json={"model": model, "messages": messages, "stream": False},
            timeout=600,
        )
        response.raise_for_status()
        content = response.json().get("message", {}).get("content", "")
        if content:
            _save_assistant(conversation_id, content)
        logger.info(
            "Ollama chat complete ok | conversation_id=%s | response_chars=%s",
            conversation_id,
            len(content or ""),
        )
        return {"success": True, "response": content, "model": model, "conversation_id": conversation_id}
    except Exception as exc:
        logger.exception("Ollama chat complete failed | conversation_id=%s | error=%s", conversation_id, exc)
        return {"success": False, "error": str(exc), "conversation_id": conversation_id}


async def _ollama_stream(model: str, messages: list[dict[str, Any]], conversation_id: int):
    chunks: list[str] = []
    logger.info(
        "Ollama chat stream start | model=%s | conversation_id=%s | messages=%s",
        model,
        conversation_id,
        len(messages),
    )
    yield json.dumps({"type": "meta", "conversation_id": conversation_id}) + "\n"
    try:
        with requests.post(
            f"{OLLAMA_URL}/api/chat",
            json={"model": model, "messages": messages, "stream": True},
            stream=True,
            timeout=600,
        ) as response:
            response.raise_for_status()
            for line in response.iter_lines():
                if not line:
                    continue
                decoded = line.decode("utf-8", errors="replace")
                try:
                    token = json.loads(decoded).get("message", {}).get("content", "")
                    if token:
                        chunks.append(token)
                except Exception:
                    pass
                yield decoded + "\n"
                await asyncio.sleep(0)
    except Exception as exc:
        logger.exception("Ollama chat stream failed | conversation_id=%s | error=%s", conversation_id, exc)
        yield json.dumps({"error": str(exc)}) + "\n"
    finally:
        if chunks:
            _save_assistant(conversation_id, "".join(chunks))
            logger.info(
                "Ollama chat stream done | conversation_id=%s | response_chars=%s",
                conversation_id,
                sum(len(c) for c in chunks),
            )


def _response(data: Any, status_code: int = 200) -> JSONResponse:
    return JSONResponse(data, status_code=status_code)


def _patched_chat_html(token: str, theme: str = "dark") -> str:
    theme = "light" if theme == "light" else "dark"
    html = TEMPLATE_PATH.read_text(encoding="utf-8", errors="replace")
    patch = f"""
<script>
window.CHAT_TOKEN = {json.dumps(token)};
window.CHAT_API_PREFIX = "/api/ollama";
window.CHAT_THEME = {json.dumps(theme)};
document.documentElement.dataset.theme = window.CHAT_THEME;
document.documentElement.classList.toggle("dark", window.CHAT_THEME === "dark");
document.documentElement.classList.toggle("light", window.CHAT_THEME === "light");
const __ftthFetch = window.fetch.bind(window);
window.fetch = function(input, init) {{
  let url = (typeof input === "string") ? input : input.url;
  if (url.startsWith("/api/")) url = window.CHAT_API_PREFIX + url.slice(4);
  if (url.startsWith("/uploads/")) url = window.CHAT_API_PREFIX + url;
  if (url.startsWith(window.CHAT_API_PREFIX)) {{
    const sep = url.includes("?") ? "&" : "?";
    url = url + sep + "_t=" + encodeURIComponent(window.CHAT_TOKEN || "");
  }}
  return __ftthFetch(url, init);
}};
</script>
<style>
html[data-theme="light"] body {{ background:#f8fafc !important; color:#0f172a !important; }}
html[data-theme="light"] .chat-app,
html[data-theme="light"] main,
html[data-theme="light"] aside,
html[data-theme="light"] section {{
  color-scheme: light;
}}
html[data-theme="dark"] body {{ background:#020617 !important; color:#e2e8f0 !important; color-scheme: dark; }}
</style>
"""
    html = html.replace("<script>\n//", patch + "\n<script>\n//", 1)
    html = html.replace("`/uploads/${apath}`", "`/api/ollama/uploads/${apath}?_t=${CHAT_TOKEN}`")
    return html


def create_ollama_chat_router(auth_dependency: Callable[..., str]) -> APIRouter:
    router = APIRouter(tags=["ollama-chat"])
    auth = Depends(auth_dependency)

    @router.get("/ui", response_class=HTMLResponse)
    async def chat_ui(request: Request, _user: str = auth):
        token = request.query_params.get("_t") or request.query_params.get("token") or ""
        theme = request.query_params.get("theme") or "dark"
        return HTMLResponse(_patched_chat_html(token, theme), headers={"Cache-Control": "no-cache, no-store"})

    @router.get("/health")
    @router.get("/v1/health")
    async def health(_user: str = auth):
        try:
            result = await asyncio.to_thread(requests.get, f"{OLLAMA_URL}/api/tags", timeout=5)
            return {"status": "ok", "ollama": "reachable", "models": len(result.json().get("models", []))}
        except Exception as exc:
            return JSONResponse({"status": "degraded", "ollama": str(exc)}, status_code=503)

    @router.get("/")
    @router.get("/docs", response_class=HTMLResponse)
    @router.get("/v1/")
    @router.get("/v1/docs", response_class=HTMLResponse)
    async def docs(_user: str = auth):
        return HTMLResponse(
            """
            <!doctype html>
            <html><head><title>Ollama Chat API</title>
            <style>body{font-family:system-ui;background:#0f172a;color:#e2e8f0;padding:32px;line-height:1.5}
            code{background:#1e293b;padding:2px 6px;border-radius:4px}li{margin:8px 0}</style></head>
            <body><h1>Ollama Chat API</h1>
            <p>Integrated into the FTTH API server. No separate Flask/nginx port is required.</p>
            <ul>
              <li><code>GET /api/ollama/health</code></li>
              <li><code>GET /api/ollama/models</code></li>
              <li><code>POST /api/ollama/chat</code> multipart chat for the embedded UI</li>
              <li><code>POST /api/ollama/v1/chat</code> JSON or multipart API chat</li>
              <li><code>GET/POST /api/ollama/conversations</code></li>
              <li><code>GET/PATCH/DELETE /api/ollama/conversations/{id}</code></li>
            </ul>
            </body></html>
            """
        )

    @router.get("/models")
    @router.get("/v1/models")
    async def models(_user: str = auth):
        try:
            result = await asyncio.to_thread(requests.get, f"{OLLAMA_URL}/api/tags", timeout=10)
            return result.json()
        except Exception as exc:
            raise HTTPException(status_code=500, detail=str(exc))

    @router.get("/conversations")
    @router.get("/v1/conversations")
    async def conversations(_user: str = auth):
        with _db() as conn:
            rows = conn.execute(
                """
                SELECT c.*,
                       (SELECT content FROM messages WHERE conversation_id=c.id
                        ORDER BY id DESC LIMIT 1) AS last_message
                FROM conversations c
                WHERE c.owner=?
                ORDER BY c.updated_at DESC
                """,
                (_user,),
            ).fetchall()
        return [dict(row) for row in rows]

    @router.post("/conversations")
    @router.post("/v1/conversations")
    async def create_conversation(request: Request, _user: str = auth):
        data = await request.json()
        with _db() as conn:
            cur = conn.execute(
                "INSERT INTO conversations (owner,title,model) VALUES (?,?,?)",
                (_user, data.get("title", "New Chat"), data.get("model", "")),
            )
            row = conn.execute("SELECT * FROM conversations WHERE id=?", (cur.lastrowid,)).fetchone()
        return _response(dict(row), status_code=201)

    @router.get("/conversations/{conversation_id}")
    @router.get("/v1/conversations/{conversation_id}")
    async def get_conversation(conversation_id: int, _user: str = auth):
        with _db() as conn:
            conv = _require_owned_conversation(conn, conversation_id, _user)
            rows = conn.execute(
                "SELECT * FROM messages WHERE conversation_id=? ORDER BY id",
                (conversation_id,),
            ).fetchall()
        return {"conversation": dict(conv), "messages": [dict(row) for row in rows]}

    @router.patch("/conversations/{conversation_id}")
    @router.patch("/v1/conversations/{conversation_id}")
    async def patch_conversation(conversation_id: int, request: Request, _user: str = auth):
        data = await request.json()
        parts: list[str] = []
        values: list[Any] = []
        for field in ("title", "model"):
            if field in data:
                parts.append(f"{field}=?")
                values.append(data[field])
        if not parts:
            raise HTTPException(status_code=400, detail="nothing to update")
        parts.append("updated_at=datetime('now')")
        values.append(conversation_id)
        values.append(_user)
        with _db() as conn:
            conn.execute(f"UPDATE conversations SET {','.join(parts)} WHERE id=? AND owner=?", values)
            row = _require_owned_conversation(conn, conversation_id, _user)
        return dict(row)

    @router.delete("/conversations/{conversation_id}")
    @router.delete("/v1/conversations/{conversation_id}")
    async def delete_conversation(conversation_id: int, _user: str = auth):
        with _db() as conn:
            _require_owned_conversation(conn, conversation_id, _user)
            conn.execute("DELETE FROM conversations WHERE id=? AND owner=?", (conversation_id, _user))
        return {"ok": True}

    async def _chat_impl(
        model: str,
        message: str,
        conversation_id: int | None,
        stream: bool,
        file: UploadFile | None,
        owner: str,
    ):
        if not model:
            raise HTTPException(status_code=400, detail="model is required")
        if not message and not file:
            raise HTTPException(status_code=400, detail="message or file is required")

        attachment_name = attachment_type = attachment_path = None
        content_kind = None
        content_data = None
        if file and file.filename:
            attachment_path, attachment_name, attachment_type = await _save_upload(file)
            content_kind, content_data = await asyncio.to_thread(_extract, attachment_path, attachment_name)

        if not conversation_id:
            conversation_id = await asyncio.to_thread(_make_conversation, owner, model)
        else:
            with _db() as conn:
                _require_owned_conversation(conn, conversation_id, owner)

        if content_kind == "image":
            db_text = message
        elif content_kind == "text" and content_data:
            db_text = f"{message}\n\n--- Attached: {attachment_name} ---\n{content_data}".strip()
        else:
            db_text = message

        await asyncio.to_thread(
            _save_user,
            conversation_id,
            db_text,
            attachment_name,
            attachment_type,
            attachment_path,
        )
        ollama_messages = await asyncio.to_thread(_load_history, conversation_id)
        if content_kind == "image" and content_data:
            ollama_messages[-1]["content"] = message or "Analyze this image."
            ollama_messages[-1].setdefault("images", [content_data])

        if stream:
            return StreamingResponse(
                _ollama_stream(model, ollama_messages, conversation_id),
                media_type="application/x-ndjson",
                headers={"X-Conversation-Id": str(conversation_id), "Cache-Control": "no-cache"},
            )

        result = await asyncio.to_thread(_ollama_complete, model, ollama_messages, conversation_id)
        if attachment_name:
            result["attachment"] = {
                "name": attachment_name,
                "type": attachment_type,
                "url": f"/api/ollama/uploads/{attachment_path}",
            }
        return result

    @router.post("/chat")
    async def chat_multipart(
        model: str = Form(""),
        message: str = Form(""),
        conversation_id: int | None = Form(None),
        stream: bool = Form(True),
        file: UploadFile | None = File(None),
        _user: str = auth,
    ):
        return await _chat_impl(model, message, conversation_id, stream, file, _user)

    @router.post("/public/extract-text")
    @router.post("/v1/public/extract-text")
    async def public_extract_text(
        model: str = Form(""),
        message: str = Form(""),
        stream: bool = Form(False),
        file: UploadFile | None = File(None),
    ):
        return await _chat_impl(model, message, None, stream, file, "public_extract")

    @router.post("/v1/chat")
    async def chat_json_or_multipart(request: Request, _user: str = auth):
        content_type = request.headers.get("content-type", "")
        if content_type.startswith("multipart/form-data"):
            form = await request.form()
            upload = form.get("file")
            upload_file = upload if hasattr(upload, "filename") and hasattr(upload, "read") else None
            return await _chat_impl(
                str(form.get("model") or ""),
                str(form.get("message") or ""),
                int(form["conversation_id"]) if form.get("conversation_id") else None,
                str(form.get("stream", "false")).lower() in {"true", "1", "yes"},
                upload_file,
                _user,
            )
        data = await request.json()
        return await _chat_impl(
            str(data.get("model") or ""),
            str(data.get("message") or ""),
            int(data["conversation_id"]) if data.get("conversation_id") else None,
            bool(data.get("stream", False)),
            None,
            _user,
        )

    @router.get("/uploads/{filename}")
    async def upload_file(filename: str, _user: str = auth):
        safe = _safe_name(filename)
        path = UPLOAD_DIR / safe
        if not path.exists() or not path.is_file():
            raise HTTPException(status_code=404, detail="file not found")
        return FileResponse(path)

    return router
