# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# Ollama Network Chat Server
#   - Web UI + file attachments (image / text / PDF / DOCX / XLSX)
#   - REST API  /api/v1/   (JSON + multipart, stream or not)
#   - SQLite history
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
from flask import (
    Flask, request, jsonify, render_template,
    Response, stream_with_context, send_from_directory,
)
import requests, json, sqlite3, uuid, base64, mimetypes
from pathlib import Path
from werkzeug.utils import secure_filename

app = Flask(__name__)

OLLAMA   = "http://127.0.0.1:11434"
DB_PATH  = Path(__file__).parent / "chat_history.db"
UP_DIR   = Path(__file__).parent / "uploads"
MAX_TEXT = 80_000  # chars fed to LLM from text files

UP_DIR.mkdir(exist_ok=True)

IMG  = {"png","jpg","jpeg","gif","webp","bmp","tiff","ico"}
TXT  = {"txt","md","py","js","ts","jsx","tsx","html","css","json","yaml","yml",
        "xml","sh","bat","ps1","c","cpp","h","hpp","java","go","rs","rb",
        "php","sql","r","swift","kt","scala","lua","toml","ini","conf","log",
        "env","dockerfile","makefile","tf","proto","vue","svelte","tex","csv"}
PDF  = {"pdf"}
DOCX = {"docx","doc"}
XLSX = {"xlsx","xls"}


# ─── Database ──────────────────────────────────────────────────

def db():
    c = sqlite3.connect(str(DB_PATH))
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys=ON")
    return c


def init_db():
    with db() as c:
        c.executescript("""
            CREATE TABLE IF NOT EXISTS conversations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL DEFAULT 'New Chat',
                model TEXT NOT NULL DEFAULT '',
                created_at TEXT DEFAULT (datetime('now')),
                updated_at TEXT DEFAULT (datetime('now'))
            );
            CREATE TABLE IF NOT EXISTS messages (
                id              INTEGER PRIMARY KEY AUTOINCREMENT,
                conversation_id INTEGER NOT NULL,
                role            TEXT    NOT NULL,
                content         TEXT    NOT NULL,
                attachment_name TEXT,
                attachment_type TEXT,
                attachment_path TEXT,
                created_at      TEXT DEFAULT (datetime('now')),
                FOREIGN KEY (conversation_id) REFERENCES conversations(id) ON DELETE CASCADE
            );
        """)
        # migrate older DBs
        have = {r[1] for r in c.execute("PRAGMA table_info(messages)")}
        for col in ("attachment_name", "attachment_type", "attachment_path"):
            if col not in have:
                c.execute(f"ALTER TABLE messages ADD COLUMN {col} TEXT")


# ─── File helpers ───────────────────────────────────────────────

def fext(n): return n.rsplit(".", 1)[-1].lower() if "." in n else ""

def fcat(n):
    e = fext(n)
    if e in IMG:  return "image"
    if e in TXT:  return "text"
    if e in PDF:  return "pdf"
    if e in DOCX: return "docx"
    if e in XLSX: return "spreadsheet"
    return "binary"


def save_upload(fo):
    """Save FileStorage → uploads/. Returns (stored_name, original_name, category)."""
    orig = secure_filename(fo.filename) or "file"
    e    = fext(orig)
    name = f"{uuid.uuid4().hex}.{e}" if e else uuid.uuid4().hex
    fo.save(str(UP_DIR / name))
    return name, orig, fcat(orig)


def extract(stored, original):
    """
    Returns (content_kind, data_string).
    content_kind: 'image' | 'text' | 'binary'
    """
    cat  = fcat(original)
    path = UP_DIR / stored

    if cat == "image":
        return "image", base64.b64encode(path.read_bytes()).decode()

    if cat == "text":
        return "text", path.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT]

    if cat == "pdf":
        try:
            import pypdf
            r    = pypdf.PdfReader(str(path))
            text = "\n\n".join(p.extract_text() or "" for p in r.pages)
            return "text", text[:MAX_TEXT]
        except Exception as e:
            return "text", f"[PDF read error: {e}]"

    if cat == "docx":
        try:
            import docx
            d = docx.Document(str(path))
            return "text", "\n".join(p.text for p in d.paragraphs if p.text)[:MAX_TEXT]
        except Exception as e:
            return "text", f"[DOCX read error: {e}]"

    if cat == "spreadsheet":
        try:
            if fext(original) == "csv":
                return "text", path.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT]
            import openpyxl
            wb   = openpyxl.load_workbook(str(path), data_only=True)
            rows = []
            for ws in wb.worksheets:
                rows.append(f"=== {ws.title} ===")
                for row in ws.iter_rows(values_only=True):
                    rows.append("\t".join(str(v or "") for v in row))
            return "text", "\n".join(rows)[:MAX_TEXT]
        except Exception as e:
            return "text", f"[Spreadsheet read error: {e}]"

    return "binary", None


# ─── Conversation helpers ───────────────────────────────────────

def load_history(conv_id, keep_imgs=6):
    """Rebuild Ollama message list from DB. Images included for recent msgs."""
    with db() as c:
        rows = c.execute(
            "SELECT * FROM messages WHERE conversation_id=? ORDER BY id", (conv_id,)
        ).fetchall()
    total = len(rows)
    out   = []
    for i, r in enumerate(rows):
        content = r["content"]
        msg = {"role": r["role"], "content": content or ""}
        recent = (total - i) <= keep_imgs
        if r["attachment_type"] == "image" and r["attachment_path"] and recent:
            p = UP_DIR / r["attachment_path"]
            if p.exists():
                msg["content"] = content or "Analyze this image."
                msg["images"]  = [base64.b64encode(p.read_bytes()).decode()]
            else:
                msg["content"] += f"\n[Image: {r['attachment_name']} – file not found]"
        out.append(msg)
    return out


def make_conv(model, title="New Chat"):
    with db() as c:
        cur = c.execute("INSERT INTO conversations (title,model) VALUES (?,?)", (title, model))
        return cur.lastrowid


def save_user(conv_id, text, a_name=None, a_type=None, a_path=None):
    with db() as c:
        c.execute(
            "INSERT INTO messages (conversation_id,role,content,attachment_name,attachment_type,attachment_path)"
            " VALUES (?,'user',?,?,?,?)",
            (conv_id, text, a_name, a_type, a_path),
        )
        cnt = c.execute(
            "SELECT COUNT(*) FROM messages WHERE conversation_id=? AND role='user'", (conv_id,)
        ).fetchone()[0]
        if cnt == 1:
            raw   = text or a_name or "New Chat"
            title = raw[:60] + ("…" if len(raw) > 60 else "")
            c.execute("UPDATE conversations SET title=?,updated_at=datetime('now') WHERE id=?",
                      (title, conv_id))
        else:
            c.execute("UPDATE conversations SET updated_at=datetime('now') WHERE id=?", (conv_id,))


def save_assistant(conv_id, text):
    with db() as c:
        c.execute(
            "INSERT INTO messages (conversation_id,role,content) VALUES (?,'assistant',?)",
            (conv_id, text),
        )
        c.execute("UPDATE conversations SET updated_at=datetime('now') WHERE id=?", (conv_id,))


# ─── Chat engine ────────────────────────────────────────────────

def do_stream(model, messages, conv_id):
    chunks = []

    def gen():
        yield json.dumps({"type": "meta", "conversation_id": conv_id}) + "\n"
        try:
            with requests.post(
                f"{OLLAMA}/api/chat",
                json={"model": model, "messages": messages, "stream": True},
                stream=True, timeout=600,
            ) as r:
                for line in r.iter_lines():
                    if line:
                        dec = line.decode()
                        try:
                            tok = json.loads(dec).get("message", {}).get("content", "")
                            if tok:
                                chunks.append(tok)
                        except Exception:
                            pass
                        yield dec + "\n"
        except Exception as e:
            yield json.dumps({"error": str(e)}) + "\n"
        finally:
            if chunks:
                save_assistant(conv_id, "".join(chunks))

    return Response(
        stream_with_context(gen()),
        mimetype="application/x-ndjson",
        headers={
            "X-Accel-Buffering": "no",
            "Cache-Control":     "no-cache",
            "X-Conversation-Id": str(conv_id),
        },
    )


def do_complete(model, messages, conv_id):
    try:
        r = requests.post(
            f"{OLLAMA}/api/chat",
            json={"model": model, "messages": messages, "stream": False},
            timeout=600,
        )
        r.raise_for_status()
        full = r.json().get("message", {}).get("content", "")
        if full:
            save_assistant(conv_id, full)
        return {"success": True, "response": full, "model": model, "conversation_id": conv_id}
    except Exception as e:
        return {"success": False, "error": str(e), "conversation_id": conv_id}


# ─── Shared request parser ──────────────────────────────────────

def parse_and_run(stream_default=True):
    """
    Parse JSON or multipart/form-data, run the LLM, return Flask response.
    Shared by internal /api/chat and external /api/v1/chat.
    """
    ct = request.content_type or ""
    if ct.startswith("multipart/form-data"):
        model   = request.form.get("model", "")
        message = request.form.get("message", "")
        conv_id = request.form.get("conversation_id", type=int)
        stream  = request.form.get("stream", str(stream_default)).lower() in ("true", "1", "yes")
        fobj    = request.files.get("file")
    else:
        data    = request.get_json(force=True, silent=True) or {}
        model   = data.get("model", "")
        message = data.get("message", "")
        conv_id = data.get("conversation_id")
        stream  = data.get("stream", stream_default)
        fobj    = None

    if not model:
        return jsonify({"success": False, "error": "model is required"}), 400
    if not message and not fobj:
        return jsonify({"success": False, "error": "message or file is required"}), 400

    # ── handle attachment
    a_name = a_type = a_path = None
    ck = cd = None                          # content_kind, content_data
    if fobj and fobj.filename:
        a_path, a_name, a_type = save_upload(fobj)
        ck, cd = extract(a_path, a_name)

    # ── ensure conversation
    if not conv_id:
        conv_id = make_conv(model)

    # ── build DB content
    if ck == "image":
        db_text = message                   # raw text only; image shown via attachment_path
    elif ck == "text" and cd:
        db_text = f"{message}\n\n--- Attached: {a_name} ---\n{cd}".strip()
    else:
        db_text = message

    save_user(conv_id, db_text, a_name, a_type, a_path)

    # ── build Ollama context
    ollama_msgs = load_history(conv_id)

    # Fix last message if it's an image (load_history handles file read,
    # but content in DB was stored as plain text)
    if ck == "image" and cd:
        last = ollama_msgs[-1]
        last["content"] = message or "Analyze this image."
        if "images" not in last:            # file might have been overwritten
            last["images"] = [cd]

    if stream:
        return do_stream(model, ollama_msgs, conv_id)
    else:
        result = do_complete(model, ollama_msgs, conv_id)
        if a_name:
            result["attachment"] = {
                "name": a_name,
                "type": a_type,
                "url":  f"/uploads/{a_path}",
            }
        return jsonify(result)


# ─── Routes: CORS for /api/v1/ ──────────────────────────────────

@app.after_request
def cors(resp):
    if request.path.startswith("/api/v1/"):
        resp.headers["Access-Control-Allow-Origin"]  = "*"
        resp.headers["Access-Control-Allow-Methods"] = "GET,POST,DELETE,PATCH,OPTIONS"
        resp.headers["Access-Control-Allow-Headers"] = "Content-Type,X-API-Key"
    return resp


@app.route("/api/v1/<path:_>", methods=["OPTIONS"])
def options_pass(_):
    return "", 204


# ─── Routes: UI & static ────────────────────────────────────────

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/uploads/<path:filename>")
def serve_upload(filename):
    return send_from_directory(str(UP_DIR), filename)


# ─── Routes: Models ─────────────────────────────────────────────

@app.route("/api/models")
@app.route("/api/v1/models")
def models():
    try:
        r = requests.get(f"{OLLAMA}/api/tags", timeout=10)
        return jsonify(r.json())
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ─── Routes: Conversations ───────────────────────────────────────

@app.route("/api/conversations",               methods=["GET"])
@app.route("/api/v1/conversations",            methods=["GET"])
def list_convs():
    with db() as c:
        rows = c.execute("""
            SELECT c.*,
                   (SELECT content FROM messages WHERE conversation_id=c.id
                    ORDER BY id DESC LIMIT 1) AS last_message
            FROM conversations c ORDER BY c.updated_at DESC
        """).fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/api/conversations",               methods=["POST"])
@app.route("/api/v1/conversations",            methods=["POST"])
def create_conv():
    d = request.get_json(force=True, silent=True) or {}
    with db() as c:
        cur = c.execute(
            "INSERT INTO conversations (title,model) VALUES (?,?)",
            (d.get("title","New Chat"), d.get("model","")),
        )
        row = c.execute("SELECT * FROM conversations WHERE id=?", (cur.lastrowid,)).fetchone()
    return jsonify(dict(row)), 201


@app.route("/api/conversations/<int:cid>",     methods=["GET"])
@app.route("/api/v1/conversations/<int:cid>",  methods=["GET"])
def get_conv(cid):
    with db() as c:
        conv = c.execute("SELECT * FROM conversations WHERE id=?", (cid,)).fetchone()
        if not conv:
            return jsonify({"error":"not found"}), 404
        msgs = c.execute(
            "SELECT * FROM messages WHERE conversation_id=? ORDER BY id", (cid,)
        ).fetchall()
    return jsonify({"conversation": dict(conv), "messages": [dict(m) for m in msgs]})


@app.route("/api/conversations/<int:cid>",     methods=["PATCH"])
@app.route("/api/v1/conversations/<int:cid>",  methods=["PATCH"])
def patch_conv(cid):
    d = request.get_json(force=True, silent=True) or {}
    parts, vals = [], []
    for f in ("title","model"):
        if f in d:
            parts.append(f"{f}=?"); vals.append(d[f])
    if not parts:
        return jsonify({"error":"nothing to update"}), 400
    parts.append("updated_at=datetime('now')"); vals.append(cid)
    with db() as c:
        c.execute(f"UPDATE conversations SET {','.join(parts)} WHERE id=?", vals)
        row = c.execute("SELECT * FROM conversations WHERE id=?", (cid,)).fetchone()
    return jsonify(dict(row))


@app.route("/api/conversations/<int:cid>",     methods=["DELETE"])
@app.route("/api/v1/conversations/<int:cid>",  methods=["DELETE"])
def del_conv(cid):
    with db() as c:
        c.execute("DELETE FROM conversations WHERE id=?", (cid,))
    return jsonify({"ok": True})


# ─── Routes: Internal chat (frontend, always streaming) ─────────

@app.route("/api/chat", methods=["POST"])
def internal_chat():
    return parse_and_run(stream_default=True)


# ─── Routes: External API v1 ────────────────────────────────────

@app.route("/api/v1/chat", methods=["POST"])
def v1_chat():
    return parse_and_run(stream_default=False)


@app.route("/api/v1/health")
def health():
    try:
        r = requests.get(f"{OLLAMA}/api/tags", timeout=5)
        models_n = len(r.json().get("models", []))
        return jsonify({"status": "ok", "ollama": "reachable", "models": models_n})
    except Exception as e:
        return jsonify({"status": "degraded", "ollama": str(e)}), 503


@app.route("/api/v1/")
@app.route("/api/v1/docs")
def v1_docs():
    return render_template("api_docs.html")


# ─── Entry point ────────────────────────────────────────────────

if __name__ == "__main__":
    import os, socket
    init_db()
    host = os.environ.get("OLLAMA_CHAT_HOST", "0.0.0.0")
    port = int(os.environ.get("OLLAMA_CHAT_PORT", "5000"))
    try:
        ip = socket.gethostbyname(socket.gethostname())
    except Exception:
        ip = "unknown"
    print(f"\n  Ollama Network Chat")
    print(f"  UI:      http://localhost:{port}")
    print(f"  Network: http://{ip}:{port}")
    print(f"  API:     http://{ip}:{port}/api/v1/\n")
    app.run(host=host, port=port, debug=False, threaded=True)
